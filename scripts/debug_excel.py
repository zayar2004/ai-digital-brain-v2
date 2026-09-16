"""Debug Excel structure detection."""
from __future__ import annotations
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook
from app.ingestion.excel_reader import _find_header_row, _normalize_header


def main() -> int:
    files = sorted(
        Path("uploads/excel").glob("*.xlsx"),
        key=lambda p: p.stat().st_mtime, reverse=True,
    )
    if not files:
        print("No uploads found.")
        return 1

    f = files[0]
    print(f"→ Debugging: {f.name}\n")

    wb = load_workbook(f, data_only=True, read_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"── Sheet: {name} ──")
        print(f"   max_row={ws.max_row} max_col={ws.max_column}")

        hdr_row, hdrs = _find_header_row(ws)
        print(f"   Detected header row: {hdr_row}")
        print(f"   Normalized headers:  {hdrs}")

        # Show first 10 raw rows
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i > 10:
                break
            print(f"   Row {i}: {list(row)}")

        # Show raw normalized values per cell
        if ws.max_row >= 1:
            raw = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
            normalized = [_normalize_header(str(x) if x else None) for x in raw]
            print(f"   Row1 raw:        {raw}")
            print(f"   Row1 normalized: {normalized}")
        print()
    wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
