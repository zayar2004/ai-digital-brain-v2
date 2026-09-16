"""
Flexible Excel reader.

- Recognizes known headers (machine/code/unit/...)
- Falls back to generic mode: treats every row as a text line
  when no known header is found.
Preserves: source_sheet, source_row.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

_HEADER_ALIASES: dict[str, str] = {
    "machine name": "machine_name",
    "machinename": "machine_name",
    "machine": "machine_name",
    "game name": "machine_name",
    "game": "machine_name",
    "name": "machine_name",
    "title": "machine_name",
    "item": "machine_name",
    "ဂိမ်း": "machine_name",

    "model": "model",
    "modelno": "model",
    "model no": "model",

    "unit": "unit",
    "position": "unit",
    "pos": "unit",
    "no": "unit",
    "number": "unit",
    "no.": "no",

    "machine code": "code",
    "machinecode": "code",
    "code": "code",
    "keycode": "code",
    "key": "code",
    "code no": "code",
    "game code": "code",

    "alias": "alias",
    "abbreviation": "alias",
    "abbr": "alias",
    "keyword": "alias",
    "short": "alias",

    "arrival date": "arrival_date",
    "arrivaldate": "arrival_date",
    "arrival": "arrival_date",
    "date": "arrival_date",

    "location": "location",
    "shop": "location",
    "section": "location",

    # Error-knowledge aliases
    "error message": "error_message",
    "error message (english)": "error_message",
    "error": "error_message",
    "error code": "error_code",
    "မြန်မာဘာသာပြန်": "myanmar_text",
    "myanmar translation": "myanmar_text",
    "possible cause": "cause",
    "ဖြစ်နိုင်သောအကြောင်း / စစ်ဆေးရန်": "cause",
    "ဖြစ်နိုင်သောအကြောင်း": "cause",
    "cause": "cause",
    "check": "check",
    "troubleshooting": "solution",
    "solution": "solution",
    "fix": "solution",
    "ဖြေရှင်းနည်း အဆင့်များ": "solution",
    "ဖြေရှင်းနည်း": "solution",
    "note": "notes",
    "notes": "notes",
    "remark": "notes",
    "မှတ်ချက်": "notes",
}


def _normalize_header(raw: str | None) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip().lower()
    s = re.sub(r"\s+", " ", s)
    s = s.replace("_", " ")
    if not s:
        return None
    return _HEADER_ALIASES.get(s, s)


@dataclass
class SheetRow:
    sheet: str
    row_number: int
    values: dict[str, str] = field(default_factory=dict)
    raw: dict[str, str] = field(default_factory=dict)
    generic_text: str = ""


@dataclass
class SheetData:
    name: str
    headers: list[str]
    rows: list[SheetRow] = field(default_factory=list)
    is_generic: bool = False


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _find_header_row(ws, max_scan: int = 10) -> tuple[int | None, list[str]]:
    known_keys = {
        "machine_name", "model", "unit", "code", "alias",
        "arrival_date", "location",
        "error_message", "error_code", "myanmar_text",
        "cause", "solution", "notes",
    }
    best: tuple[int, list[str]] | None = None
    best_count = 0
    for r in range(1, min(ws.max_row, max_scan) + 1):
        raw = [_cell_to_str(ws.cell(row=r, column=c).value)
               for c in range(1, ws.max_column + 1)]
        non_empty = [x for x in raw if x]
        if len(non_empty) < 2:
            continue
        normalized = [_normalize_header(x) for x in raw]
        known = sum(1 for h in normalized if h in known_keys)
        if known > best_count:
            best_count = known
            best = (r, normalized)
            if known >= 3:
                break
    if best and best_count >= 1:
        return best
    return None, []


def read_excel(path: str | Path, *, generic_fallback: bool = True) -> list[SheetData]:
    """
    Read an Excel workbook.

    For each sheet:
      1. Try to detect known headers.
      2. If no known headers AND generic_fallback=True, treat the whole
         sheet as a generic text table (each row becomes one line).
    """
    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    results: list[SheetData] = []

    for ws in wb.worksheets:
        header_row, headers = _find_header_row(ws)

        if header_row is not None and headers:
            sheet = SheetData(name=ws.title, headers=headers, is_generic=False)
            for r in range(header_row + 1, ws.max_row + 1):
                row_values: dict[str, str] = {}
                raw_values: dict[str, str] = {}
                has_data = False
                for c in range(1, ws.max_column + 1):
                    cell = _cell_to_str(ws.cell(row=r, column=c).value)
                    header_key = headers[c - 1] if c - 1 < len(headers) else None
                    if not header_key:
                        continue
                    raw_values[header_key] = cell
                    if cell:
                        has_data = True
                        row_values[header_key] = cell
                if has_data:
                    sheet.rows.append(SheetRow(
                        sheet=ws.title, row_number=r,
                        values=row_values, raw=raw_values,
                    ))
            if sheet.rows:
                results.append(sheet)
        elif generic_fallback and ws.max_row >= 1:
            # Generic: treat every non-empty row as a text line
            sheet = SheetData(
                name=ws.title, headers=[], is_generic=True,
            )
            for r in range(1, ws.max_row + 1):
                cells = [
                    _cell_to_str(ws.cell(row=r, column=c).value)
                    for c in range(1, ws.max_column + 1)
                ]
                non_empty = [c for c in cells if c]
                if not non_empty:
                    continue
                text = " | ".join(non_empty)
                sheet.rows.append(SheetRow(
                    sheet=ws.title, row_number=r,
                    values={}, raw={}, generic_text=text,
                ))
            if sheet.rows:
                results.append(sheet)

    wb.close()
    return results
